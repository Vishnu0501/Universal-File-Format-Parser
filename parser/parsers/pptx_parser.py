import io
from typing import Iterator
from parser.record import Record, RecordType, Severity

try:
    from pptx import Presentation
    _PPTX = True
except ImportError:
    _PPTX = False


def parse_pptx(data: bytes) -> Iterator[Record]:
    if not _PPTX:
        yield Record(type=RecordType.ERROR, format="pptx", reason="python-pptx is not installed")
        return

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        yield Record(type=RecordType.ERROR, format="pptx", reason=f"Could not open PPTX: {e}", offset=0)
        return

    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        try:
            for shape in _shapes_in_order(slide):
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        texts.append(text)

            if texts:
                yield Record(
                    type=RecordType.DATA,
                    format="pptx",
                    content="\n".join(texts),
                    metadata={"slide": slide_num},
                )
            else:
                yield Record(
                    type=RecordType.WARNING,
                    format="pptx",
                    reason=f"Slide {slide_num} has no extractable text",
                    severity=Severity.LOW,
                    metadata={"slide": slide_num},
                )

        except Exception as e:
            yield Record(
                type=RecordType.ERROR,
                format="pptx",
                reason=f"Error reading slide {slide_num}: {e}",
                metadata={"slide": slide_num},
            )


def _shapes_in_order(slide) -> list:
    try:
        return sorted(slide.shapes, key=lambda s: (s.top or 0, s.left or 0))
    except Exception:
        return list(slide.shapes)