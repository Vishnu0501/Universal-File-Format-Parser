import io
from typing import Iterator, List, Any
from parser.record import Record, RecordType, Severity

try:
    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException
    _OPENPYXL = True
except ImportError:
    _OPENPYXL = False


def parse_xlsx(data: bytes) -> Iterator[Record]:
    if not _OPENPYXL:
        yield Record(type=RecordType.ERROR, format="xlsx", reason="openpyxl is not installed")
        return

    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(data),
            read_only=True,
            data_only=True,
        )
    except InvalidFileException as e:
        yield Record(type=RecordType.ERROR, format="xlsx", reason=f"Invalid XLSX file: {e}", offset=0)
        return
    except Exception as e:
        yield Record(type=RecordType.ERROR, format="xlsx", reason=f"Could not open XLSX: {e}", offset=0)
        return

    try:
        sheet_names = wb.sheetnames
    except Exception as e:
        yield Record(type=RecordType.ERROR, format="xlsx", reason=f"Could not read sheet names: {e}")
        return

    for sheet_index, sheet_name in enumerate(sheet_names):
        try:
            ws = wb[sheet_name]
            grid = _sheet_to_grid(ws)

            yield Record(
                type=RecordType.DATA,
                format="xlsx",
                content=grid,
                metadata={
                    "sheet": sheet_name,
                    "sheet_index": sheet_index,
                    "rows": len(grid),
                    "cols": len(grid[0]) if grid else 0,
                },
            )
        except Exception as e:
            yield Record(
                type=RecordType.ERROR,
                format="xlsx",
                reason=f"Error reading sheet '{sheet_name}': {e}",
                metadata={"sheet": sheet_name, "sheet_index": sheet_index},
            )

    try:
        wb.close()
    except Exception:
        pass


def _sheet_to_grid(ws) -> List[List[Any]]:
    grid = []
    for row in ws.iter_rows():
        grid.append([_cell_value(cell) for cell in row])
    return grid


def _cell_value(cell) -> Any:
    val = cell.value
    if val is None:
        return None
    if hasattr(val, "isoformat"):
        return val.isoformat()
    return val