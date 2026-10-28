import io
from typing import Iterator
from parser.record import Record, RecordType, Severity

try:
    from docx import Document
    _DOCX = True
except ImportError:
    _DOCX = False


def parse_docx(data: bytes) -> Iterator[Record]:
    if not _DOCX:
        yield Record(type=RecordType.ERROR, format="docx", reason="python-docx is not installed")
        return

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        yield Record(type=RecordType.ERROR, format="docx", reason=f"Could not open DOCX: {e}", offset=0)
        return

    paragraphs = []
    try:
        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    paragraphs.append(" | ".join(row_text))

    except Exception as e:
        yield Record(type=RecordType.ERROR, format="docx", reason=f"Error reading DOCX content: {e}")
        return

    if paragraphs:
        yield Record(
            type=RecordType.DATA,
            format="docx",
            content="\n".join(paragraphs),
            metadata={"paragraph_count": len(paragraphs)},
        )
    else:
        yield Record(
            type=RecordType.WARNING,
            format="docx",
            reason="DOCX contains no extractable text",
            severity=Severity.LOW,
        )