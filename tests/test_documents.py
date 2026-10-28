import io
import pytest
import openpyxl
from parser.record import RecordType
from parser.parsers.xlsx_parser import parse_xlsx


def data_records(gen):
    return [r for r in gen if r.type == RecordType.DATA]


def error_records(gen):
    return [r for r in gen if r.type == RecordType.ERROR]


# ── XLSX ────────────────────────────────────────────────────────────────────────

def _make_xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_xlsx_single_sheet():
    data = _make_xlsx({"Sheet1": [["name", "age"], ["Alice", 30], ["Bob", 25]]})
    rows = data_records(parse_xlsx(data))
    assert len(rows) == 1
    assert rows[0].content[0] == ["name", "age"]
    assert rows[0].content[1] == ["Alice", 30]


def test_xlsx_multi_sheet():
    data = _make_xlsx({
        "Sheet1": [["a", "b"], [1, 2]],
        "Sheet2": [["x", "y"], [3, 4]],
    })
    rows = data_records(parse_xlsx(data))
    assert len(rows) == 2
    assert rows[0].metadata["sheet"] == "Sheet1"
    assert rows[1].metadata["sheet"] == "Sheet2"


def test_xlsx_sheet_index_in_order():
    data = _make_xlsx({
        "Alpha": [[1]],
        "Beta": [[2]],
        "Gamma": [[3]],
    })
    rows = data_records(parse_xlsx(data))
    indices = [r.metadata["sheet_index"] for r in rows]
    assert indices == [0, 1, 2]


def test_xlsx_empty_cells_preserved():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    ws["C1"] = "world"
    buf = io.BytesIO()
    wb.save(buf)
    data = buf.getvalue()

    rows = data_records(parse_xlsx(data))
    grid = rows[0].content
    assert grid[0][0] == "hello"
    assert grid[0][1] is None
    assert grid[0][2] == "world"


def test_xlsx_metadata_has_shape():
    data = _make_xlsx({"Sheet1": [["a", "b", "c"], [1, 2, 3], [4, 5, 6]]})
    rows = data_records(parse_xlsx(data))
    assert rows[0].metadata["rows"] == 3
    assert rows[0].metadata["cols"] == 3


def test_xlsx_invalid_bytes_yields_error():
    errs = error_records(parse_xlsx(b"not a zip file at all"))
    assert len(errs) >= 1


def test_xlsx_date_serialized_as_iso():
    from datetime import date
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = date(2024, 6, 1)
    buf = io.BytesIO()
    wb.save(buf)
    rows = data_records(parse_xlsx(buf.getvalue()))
    val = rows[0].content[0][0]
    assert isinstance(val, str)
    assert "2024" in val


# ── PDF / DOCX / PPTX (smoke tests — require real files or mocks) ───────────────
# These are integration-level smoke tests using minimal real byte structures.

def test_pdf_parser_wrong_bytes_yields_error():
    from parser.parsers.pdf_parser import parse_pdf
    errs = error_records(parse_pdf(b"not a pdf"))
    assert len(errs) >= 1


def test_docx_parser_wrong_bytes_yields_error():
    from parser.parsers.docx_parser import parse_docx
    errs = error_records(parse_docx(b"not a docx"))
    assert len(errs) >= 1


def test_pptx_parser_wrong_bytes_yields_error():
    from parser.parsers.pptx_parser import parse_pptx
    errs = error_records(parse_pptx(b"not a pptx"))
    assert len(errs) >= 1


def test_docx_real_content():
    from docx import Document
    from parser.parsers.docx_parser import parse_docx
    doc = Document()
    doc.add_paragraph("Hello from paragraph one.")
    doc.add_paragraph("Hello from paragraph two.")
    buf = io.BytesIO()
    doc.save(buf)
    rows = data_records(parse_docx(buf.getvalue()))
    assert len(rows) == 1
    assert "Hello from paragraph one." in rows[0].content
    assert "Hello from paragraph two." in rows[0].content


def test_pptx_real_content_ordered_by_slide():
    from pptx import Presentation
    from pptx.util import Inches
    from parser.parsers.pptx_parser import parse_pptx
    prs = Presentation()
    layout = prs.slide_layouts[5]
    for i in range(3):
        slide = prs.slides.add_slide(layout)
        txBox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        txBox.text_frame.text = f"Slide {i + 1} content"
    buf = io.BytesIO()
    prs.save(buf)
    rows = data_records(parse_pptx(buf.getvalue()))
    assert len(rows) == 3
    slides = [r.metadata["slide"] for r in rows]
    assert slides == [1, 2, 3]
